from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
from decimal import Decimal
from django.core.exceptions import ValidationError
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .services import summary

def safe(value):
    if isinstance(value, str):
        # Excel XML cannot encode NUL/control characters. Keep the source unchanged.
        value = ''.join(c if c in '\t\n\r' or 0x20 <= ord(c) <= 0xD7FF or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF else '\ufffd' for c in value)
    # Spreadsheet formula injection protection: all user text remains text.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value

def workbook(sheets):
    book = Workbook()
    book.remove(book.active)
    for title, headers, rows in sheets:
        sheet = book.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([safe(v) for v in row])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='163F3C')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(16, max(len(str(c.value or '')) for c in column) + 3))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type == 'n':
                    if isinstance(cell.value, int):
                        cell.number_format = '#,##0'
                    else:
                        scale = max(2, min(8, -cell.value.as_tuple().exponent)) if isinstance(cell.value, Decimal) else 2
                        cell.number_format = '#,##0.' + '0' * scale
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    out = BytesIO()
    book.save(out)
    return out.getvalue()

def archive_bytes(order):
    common_headers = ['订单号', '客户', '业务员', '小项', '下单日期', '产品', '型号', '数量', '销售单价', '币种', '报价汇率 CNY/销售币', '销售总额', '单位成本 CNY', '采购总额 CNY', '供应商', '单价含运费', '客户运费', '客户条款', '供应商条款', '供货方式', '对外货款 CNY']
    common = [[order.number, order.customer, order.owner.get_full_name() or order.owner.username, l.number, l.date.isoformat(), l.product, l.model, l.quantity, l.price, l.currency, l.quote_fx, l.sale_total, l.cost, l.cost_total, l.supplier, str(l.freight_included), l.freight, l.customer_terms, l.supplier_terms, l.get_supply_type_display(), l.payable_total] for l in order.lines.all()]
    files = {}
    for kind, filename in [('receipt', '客户收款表.xlsx'), ('payment', '供应商付款表.xlsx')]:
        transactions, allocations = [], []
        for t in order.transactions.filter(kind=kind):
            att = f'凭证/{t.id}{__import__("pathlib").Path(t.attachment.name).suffix}' if t.attachment else ''
            transactions.append([t.id, order.number, t.date.isoformat(), t.amount, t.currency, t.effective_party, t.platform, t.fx, t.fx_source, t.equivalent, t.target_currency, t.purpose, t.reference, t.unallocated, att, t.created_by.username, t.party, t.effective_amount, '已冲销误录' if t.is_void else '有效'])
            if t.attachment:
                try:
                    with t.attachment.open('rb') as f:
                        files[att] = f.read()
                except OSError:
                    raise ValidationError(f'流水 #{t.id} 的凭证文件暂不可用，已停止导出，请核对文件备份后重试。')
            for a in t.allocations.select_related('line', 'shipment_item'):
                allocations.append([t.id, a.line.number, a.line.model, a.source_amount, t.currency, a.settlement_rate, a.target_amount, a.line.currency if kind == 'receipt' else 'CNY', a.shipment_item.shipment_id if a.shipment_item else '未发货资金池', a.reason, a.created_by.username, '；'.join(f'批次 #{c.shipment_item.shipment_id}: {c.amount}' for c in a.coverages.select_related('shipment_item'))])
        files[filename] = workbook([('流水', ['流水ID', '订单号', '日期', '原币金额', '币种', '付款公司/供应商', '平台', '折算汇率', '汇率来源', '等值金额', '等值币种', '用途', '平台流水号', '未分配原币金额', '凭证文件', '录入人', '最初录入公司', '有效原币金额', '流水状态'], transactions), ('款项分摊', ['流水ID', '小项', '型号', '分摊原币', '原币种', '结算汇率', '抵扣金额', '抵扣币种', '发货批次', '说明', '操作人', '自动结算批次（目标币金额）'], allocations), ('订单小项快照', common_headers, common)])
    shipments = []
    for s in order.shipments.all():
        for i in s.items.select_related('line'):
            shipments.append([s.id, s.date.isoformat(), i.line.number, i.line.product, i.line.model, i.quantity, i.sale_amount, i.line.currency, i.receipt_reserved, i.covered('receipt'), i.due('receipt'), i.cost_amount, i.payment_reserved, i.covered('payment'), i.due('payment'), s.address, s.get_carrier_type_display(), s.carrier, str(s.domestic_date or ''), s.domestic_tracking, s.international_tracking, s.current_freight, s.current_freight_currency, s.notes, i.internal_cost_amount])
    files['发货表.xlsx'] = workbook([('发货明细', ['批次', '日期', '小项', '产品', '型号', '数量', '本批货值', '销售币种', '确认时占用客户预付款', '目前已覆盖客户款', '本批待收', '本批对外货款 CNY', '确认时占用供应商预付款', '目前已覆盖供应商款', '本批待付 CNY', '收货地址', '运输方', '货代', '国内日期', '国内单号', '国际单号', '批次运费(同批多行勿重复汇总)', '运费币种', '备注', '本批内部成本 CNY'], shipments), ('运费批次汇总', ['批次', '当前实际运费', '币种', '状态'], [[s.pk, s.current_freight, s.current_freight_currency, '待补录' if s.current_freight is None else '已记录'] for s in order.shipments.all()]),
        ('运费历史', ['批次', '金额', '币种', '说明', '操作人', '时间'], [[s.pk, r.amount, r.currency, r.reason, r.created_by.username, r.created_at.isoformat()] for s in order.shipments.all() for r in s.freight_records.select_related('created_by')]),
        ('订单小项快照', common_headers, common)])
    files['流水更正及冲销.json'] = json.dumps([{'transaction_id': t.pk, 'original_party': t.party, 'effective_party': t.effective_party, 'original_amount': str(t.amount), 'currency': t.currency, 'effective_amount': str(t.effective_amount), 'corrections': [{'id': c.pk, 'action': c.action, 'party': c.party, 'source_delta': str(c.source_delta), 'reason': c.reason, 'actor': c.created_by.username, 'time': c.created_at.isoformat()} for c in t.corrections.select_related('created_by')]} for t in order.transactions.all()], ensure_ascii=False, indent=2).encode('utf-8')
    files['审计记录.json'] = json.dumps([{'time': a.created_at.isoformat(), 'actor': a.actor.username, 'action': a.action, 'detail': a.detail} for a in order.audits.select_related('actor')], ensure_ascii=False, indent=2).encode('utf-8')
    manifest = {'order': order.number, 'complete': summary(order)['complete'], 'files': {name: hashlib.sha256(data).hexdigest() for name, data in files.items()}}
    files['校验清单.json'] = json.dumps(manifest, ensure_ascii=False, indent=2).encode('utf-8')
    out = BytesIO()
    with ZipFile(out, 'w', ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return out.getvalue()
