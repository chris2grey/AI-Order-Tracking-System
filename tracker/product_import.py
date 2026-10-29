from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from .exports import workbook

COLUMNS = [
    ('product', '产品名称'),
    ('model', '产品型号'),
    ('quantity', '数量'),
    ('price', '销售单价'),
    ('currency', '销售币种'),
    ('supply_type', '供货方式'),
    ('supplier', '供应商名称'),
    ('cost', '单位成本CNY'),
    ('freight', '客户运费'),
]

def template_bytes():
    return workbook([
        ('产品录入', [label for key, label in COLUMNS], []),
        ('填写说明', ['项目', '要求'], [
            ['使用流程', '填写产品录入工作表 → 上传 → 页面核对统一日期、汇率和条款 → 保存全部产品。上传本身不写入订单。'],
            ['行数', '一次最多100行产品，请保留第一行列标题及工作表名称。'],
            ['供货方式', '填写 外部供应商 或 自供货。自供货不产生对外待付；内部成本仍需填写。'],
            ['销售币种', 'USD / CNY / EUR / HKD / GBP；汇率在网页统一设置中每种填写一次。'],
            ['必填', '产品名称、型号、正整数数量、正销售单价、销售币种、供货方式、非负单位成本；外部供货必填供应商。'],
            ['可空', '小项编号由系统生成，无需填写；客户运费留空为0。日期、付款条款及单价是否含运费在网页统一设置。'],
            ['安全与核对', '不支持公式、宏及合并单元格。导入先生成草稿，确认保存后才更新订单。相同型号不同补单不自动合并。'],
        ]),
    ])

def import_rows(upload):
    if not upload or Path(upload.name).suffix.lower() != '.xlsx':
        raise ValidationError('请选择下载模板填写后的 .xlsx 文件。')
    if upload.size > 5 * 1024 * 1024:
        raise ValidationError('模板文件最大 5MB。')
    data = upload.read()
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 250 or sum(e.file_size for e in entries) > 25 * 1024 * 1024:
                raise ValidationError('表格内容过大，请使用固定模板，最多100行。')
            if any('vbaproject' in e.filename.lower() for e in entries):
                raise ValidationError('不接受带宏的工作簿。')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if '产品录入' not in book.sheetnames:
                raise ValidationError('缺少“产品录入”工作表，请重新下载模板。')
            sheet = book['产品录入']
            if sheet.max_row and sheet.max_row > 101:
                raise ValidationError('单次最多导入100行，请移除多余空行及格式。')
            legacy = next(sheet.iter_rows(max_row=1, values_only=True))[0] == '小项编号（可空）'
            columns = ([('label', '小项编号（可空）')] if legacy else []) + COLUMNS
            if sheet.max_column and sheet.max_column > len(columns):
                raise ValidationError('表格列数与模板不一致。')
            rows = sheet.iter_rows(max_row=101, max_col=len(columns))
            headers = [c.value for c in next(rows)]
            if headers != [label for key, label in columns]:
                raise ValidationError('列标题或顺序与模板不一致，请下载固定模板。')
            result = []
            for row in rows:
                if any(c.data_type == 'f' for c in row):
                    raise ValidationError(f'第{row[0].row}行含公式，请粘贴为实际数值。')
                values = [c.value for c in row]
                if all(value is None or value == '' for value in values):
                    continue
                entry = {key: str(value).strip() if value is not None else '' for (key, label), value in zip(columns, values)}
                entry.pop('label', None)  # Older templates remain usable, but manual labels are ignored.
                entry['currency'] = entry['currency'].upper()
                entry['supply_type'] = {'外部供应商':'external', '自供货':'self', 'external':'external', 'self':'self'}.get(entry['supply_type'], entry['supply_type'])
                entry['freight'] = entry['freight'] or '0'
                result.append(entry)
            if not result:
                raise ValidationError('模板中没有产品，请先填写再上传。')
            return result
        finally:
            book.close()
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('无法读取此表格，请使用下载的 .xlsx 模板。') from exc
