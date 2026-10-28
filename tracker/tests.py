from datetime import date
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile
import json
import uuid
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError, PermissionDenied
from django.test import TestCase, Client
from openpyxl import load_workbook
from .models import Order, Line, Transaction, Shipment, Allocation
from .services import allocate, shipment_quote, confirm_shipment, summary
from .exports import archive_bytes

class BusinessTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user('admin', password='random-test-password', is_superuser=True)
        cls.alice = User.objects.create_user('alice', password='random-test-password')
        cls.bob = User.objects.create_user('bob', password='random-test-password')
        cls.order = Order.objects.create(number='TEST-001', customer='Test Customer', owner=cls.alice)
        cls.a = Line.objects.create(order=cls.order, label='001', product='A', model='A', quantity=5000, price=10, cost=40, supplier='Factory A')
        cls.b = Line.objects.create(order=cls.order, label='002', product='B', model='B', quantity=1000, price=20, cost=80, supplier='Factory B')
        cls.c = Line.objects.create(order=cls.order, label='003', product='C', model='C', quantity=1000, price=30, cost=120, supplier='Factory C')
        cls.other = Order.objects.create(number='TEST-002', customer='Hidden Customer', owner=cls.bob)
        cls.hidden = Line.objects.create(order=cls.other,label='001',product='Hidden',model='Hidden',quantity=10,price=1,cost=1,supplier='Hidden Factory')

    def txn(self, amount, kind='receipt', line=None, currency=None, order=None, party=None, fx=1):
        return Transaction.objects.create(order=order or self.order, kind=kind, amount=amount, currency=currency or ('USD' if kind=='receipt' else 'CNY'), party=party or ((line or self.a).supplier if kind=='payment' else 'Test Company'),platform='Bank',fx=fx,purpose='Test',reference=uuid.uuid4().hex,created_by=self.admin)

    def fund(self):
        for l, amount, kind in [(self.a,15000,'receipt'),(self.c,30000,'receipt'),(self.a,100000,'payment'),(self.b,16000,'payment'),(self.c,120000,'payment')]:
            t=self.txn(amount,kind,l)
            allocate(self.admin,t.id,l.id,amount,1,'Initial')

    def ship(self, quantities, token=None):
        return confirm_shipment(self.admin,self.order.id,quantities,{'date':date.today(),'address':'Address','carrier_type':'own','carrier':'Carrier','freight':0},token or uuid.uuid4())

    def test_abc_expected_balances(self):
        self.fund()
        quotes=[shipment_quote(l,1000) for l in [self.a,self.b,self.c]]
        self.assertEqual([q['receipt']['due'] for q in quotes],[Decimal('7000'),Decimal('20000'),Decimal('0')])
        self.assertEqual(sum(q['receipt']['due'] for q in quotes),Decimal('27000'))
        self.assertEqual([q['payment']['due'] for q in quotes],[Decimal('20000'),Decimal('64000'),Decimal('0')])

    def test_same_model_supplement_stays_independent(self):
        b2=Line.objects.create(order=self.order,label='004',product='B',model='B',quantity=2000,price=22,cost=85,supplier='Factory B')
        t=self.txn(1000)
        allocate(self.admin,t.id,b2.id,1000,1,'Second B')
        self.assertEqual(shipment_quote(self.b,1000)['receipt']['covered'],0)
        self.assertEqual(shipment_quote(b2,1000)['receipt']['covered'],500)

    def test_overallocation_rejected(self):
        t=self.txn(100)
        allocate(self.admin,t.id,self.a.id,80,1,'First')
        with self.assertRaises(ValidationError): allocate(self.admin,t.id,self.b.id,21,1,'Too much')
        self.assertEqual(t.allocations.count(),1)

    def test_line_overpayment_rejected(self):
        t=self.txn(60000)
        with self.assertRaises(ValidationError): allocate(self.admin,t.id,self.a.id,50001,1,'Too much')

    def test_supplier_mismatch_rejected(self):
        t=self.txn(100,'payment',self.a)
        with self.assertRaises(ValidationError): allocate(self.admin,t.id,self.b.id,100,1,'Wrong factory')

    def test_cross_order_allocation_rejected(self):
        t=self.txn(10)
        with self.assertRaises(ValidationError): allocate(self.admin,t.id,self.hidden.id,10,1,'Wrong order')

    def test_salesperson_cannot_confirm_finance(self):
        t=self.txn(100)
        with self.assertRaises(PermissionDenied): allocate(self.alice,t.id,self.a.id,100,1,'No permission')

    def test_historical_batch_reservations_not_reused(self):
        self.fund()
        first=self.ship({self.a.id:1000}).items.get()
        second=self.ship({self.a.id:1000}).items.get()
        self.assertEqual(first.receipt_reserved,3000)
        self.assertEqual(second.receipt_reserved,3000)
        self.assertEqual(shipment_quote(self.a,3000)['receipt']['covered'],9000)

    def test_batch_tail_payment_only_funds_target_batch(self):
        self.fund()
        item=self.ship({self.a.id:1000}).items.get()
        t=self.txn(7000)
        allocate(self.admin,t.id,self.a.id,7000,1,'Tail',item.id)
        self.assertEqual(item.due('receipt'),0)
        self.assertEqual(shipment_quote(self.a,1000)['receipt']['covered'],3000)

    def test_unassigned_tail_requires_batch(self):
        self.fund()
        self.ship({self.a.id:1000})
        t=self.txn(7000)
        with self.assertRaises(ValidationError): allocate(self.admin,t.id,self.a.id,7000,1,'Must target batch')

    def test_shipping_over_quantity_rejected_atomically(self):
        with self.assertRaises(ValidationError): self.ship({self.a.id:100,self.b.id:1001})
        self.assertEqual(Shipment.objects.count(),0)

    def test_cross_order_shipping_rejected(self):
        with self.assertRaises(ValidationError): self.ship({self.hidden.id:1})

    def test_shipment_confirmation_idempotent(self):
        token=uuid.uuid4()
        a=self.ship({self.a.id:10},token)
        b=self.ship({self.a.id:10},token)
        self.assertEqual(a.id,b.id)
        self.assertEqual(Shipment.objects.count(),1)

    def test_rounding_final_batch_reconciles(self):
        self.a.quantity=3;self.a.price=Decimal('0.3333');self.a.cost=Decimal('0.3333');self.a.save()
        t=self.txn('0.50');allocate(self.admin,t.id,self.a.id,'0.50',1,'Half')
        for _ in range(3): self.ship({self.a.id:1})
        items=list(self.a.shipment_items.all())
        self.assertEqual(sum(i.sale_amount for i in items),Decimal('1.00'))
        self.assertEqual(sum(i.receipt_reserved for i in items),Decimal('0.50'))

    def test_cross_currency_uses_settlement_rate_not_reporting_rate(self):
        t=self.txn(700,'receipt',currency='CNY',fx=Decimal('0.14'))
        a=allocate(self.admin,t.id,self.a.id,700,Decimal('0.14285714'),'Agreed rate')
        self.assertEqual(a.target_amount,Decimal('100.00'))
        self.assertEqual(t.equivalent,Decimal('98.00'))

    def test_same_currency_rate_must_be_one(self):
        t=self.txn(100)
        with self.assertRaises(ValidationError): allocate(self.admin,t.id,self.a.id,100,2,'Invalid')

    def test_unallocated_funds_do_not_count_as_received(self):
        self.txn(1000)
        s=summary(self.order)
        self.assertEqual(s['currencies']['USD']['receipt'],0)
        self.assertEqual(len(s['unallocated']),1)

    def test_shipping_freight_not_double_counted(self):
        self.a.freight=100;self.a.freight_included=True;self.a.save()
        self.assertEqual(self.a.sale_total,50000)
        self.a.freight_included=False;self.a.save()
        self.assertEqual(self.a.sale_total,50100)

    def test_all_pages_render_for_admin(self):
        self.client.force_login(self.admin)
        t=self.txn(100)
        for url in ['/', '/orders/',f'/orders/{self.order.pk}/','/orders/new/',f'/orders/{self.order.pk}/lines/new/',f'/orders/{self.order.pk}/transactions/receipt/new/',f'/orders/{self.order.pk}/transactions/payment/new/',f'/transactions/{t.pk}/allocate/',f'/orders/{self.order.pk}/shipments/new/','/finance/','/shipments/','/archives/','/team/']:
            with self.subTest(url=url): self.assertEqual(self.client.get(url).status_code,200)

    def test_salesperson_data_isolation(self):
        self.client.force_login(self.alice)
        t=self.txn(10,order=self.other)
        for url in [f'/orders/{self.other.pk}/',f'/orders/{self.other.pk}/export/',f'/orders/{self.other.pk}/lines/new/',f'/orders/{self.other.pk}/shipments/new/',f'/transactions/{t.pk}/attachment/',f'/transactions/{t.pk}/allocate/']:
            with self.subTest(url=url): self.assertEqual(self.client.get(url).status_code,404)
        for url in ['/','/orders/','/finance/','/archives/','/shipments/']:
            response=self.client.get(url)
            self.assertNotContains(response,'Hidden Customer')

    def test_salesperson_finance_and_team_write_forbidden(self):
        self.client.force_login(self.alice)
        t=self.txn(10)
        for url in [f'/orders/{self.order.pk}/transactions/receipt/new/',f'/transactions/{t.pk}/allocate/','/team/']:
            self.assertEqual(self.client.get(url).status_code,403)

    def test_export_contains_three_workbooks_and_manifest(self):
        self.fund();self.ship({self.a.id:1000})
        self.a.product='=HYPERLINK("https://example.com")';self.a.save()
        with ZipFile(BytesIO(archive_bytes(self.order))) as z:
            for name in ['客户收款表.xlsx','供应商付款表.xlsx','发货表.xlsx']:
                book=load_workbook(BytesIO(z.read(name)))
                self.assertIn('订单小项快照',book.sheetnames)
                self.assertNotEqual(book['订单小项快照']['F2'].data_type,'f')
            manifest=json.loads(z.read('校验清单.json'))
            import hashlib
            for name,digest in manifest['files'].items(): self.assertEqual(hashlib.sha256(z.read(name)).hexdigest(),digest)

    def test_incomplete_order_cannot_archive(self):
        self.client.force_login(self.admin)
        self.client.post(f'/orders/{self.order.pk}/lock/')
        self.order.refresh_from_db()
        self.assertEqual(self.order.status,'active')
        self.assertEqual(self.order.archives.count(),0)

    def test_archive_lock_blocks_later_business_mutations(self):
        for line in [self.a,self.b,self.c]:
            for kind,amount in [('receipt',line.sale_total),('payment',line.cost_total)]:
                t=self.txn(amount,kind,line);allocate(self.admin,t.id,line.id,amount,1,'Full')
        self.ship({self.a.id:5000,self.b.id:1000,self.c.id:1000})
        self.assertTrue(summary(self.order)['complete'])
        self.client.force_login(self.admin)
        self.client.post(f'/orders/{self.order.pk}/lock/')
        self.order.refresh_from_db()
        self.assertEqual(self.order.status,'locked')
        self.assertEqual(self.order.archives.count(),1)
        with self.assertRaises(ValidationError): self.ship({self.a.id:1})
        self.assertEqual(self.client.get(f'/archives/{self.order.archives.get().pk}/download/').status_code,200)

    def test_csrf_required(self):
        client=Client(enforce_csrf_checks=True);client.force_login(self.admin)
        self.assertEqual(client.post(f'/orders/{self.order.pk}/lock/').status_code,403)

    def test_anonymous_redirected(self):
        for url in ['/','/finance/',f'/orders/{self.order.pk}/export/']:
            self.assertEqual(self.client.get(url).status_code,302)

    def test_upload_rejects_mislabeled_file(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .forms import TransactionForm
        f=TransactionForm(files={'attachment':SimpleUploadedFile('fake.pdf',b'<html>attack</html>')})
        f.is_valid()
        self.assertIn('attachment',f.errors)

    def test_quantity_only_shipment_preview(self):
        self.fund()
        self.client.force_login(self.admin)
        response=self.client.post(f'/orders/{self.order.pk}/shipments/new/', {'action':'preview','token':str(uuid.uuid4()),f'qty_{self.a.id}':'1000'})
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'7,000.00 USD')
        self.assertNotContains(response,'这个字段是必填项')
        self.assertEqual(Shipment.objects.count(),0)

    def test_browser_form_workflow_end_to_end(self):
        self.client.force_login(self.admin)
        response=self.client.post('/orders/new/',{'number':'FORM-001','customer':'Form Customer','owner':self.alice.pk,'date':'2026-10-23','notes':'Test'})
        self.assertEqual(response.status_code,302)
        o=Order.objects.get(number='FORM-001')
        response=self.client.post(f'/orders/{o.pk}/lines/new/',{'label':'001','product':'Module','model':'FORM-A','date':'2026-10-23','quantity':100,'price':'10','currency':'USD','quote_fx':'7','cost':'40','supplier':'Form Factory','freight_included':'on','freight':'0','customer_terms':'Before shipping','supplier_terms':'Before shipping'})
        self.assertEqual(response.status_code,302)
        l=o.lines.get()
        for kind,amount,currency,party in [('receipt',1000,'USD','Form Company'),('payment',4000,'CNY','Form Factory')]:
            response=self.client.post(f'/orders/{o.pk}/transactions/{kind}/new/',{'date':'2026-10-23','amount':amount,'currency':currency,'party':party,'platform':'Bank','fx':'1','fx_source':'Manual','purpose':'Full','reference':f'FORM-{kind}'})
            self.assertEqual(response.status_code,302)
            t=o.transactions.get(kind=kind)
            response=self.client.post(f'/transactions/{t.pk}/allocate/',{'line':l.pk,'amount':amount,'rate':'1','reason':'Full','shipment_item':''})
            self.assertEqual(response.status_code,302)
        response=self.client.post(f'/orders/{o.pk}/shipments/new/',{'action':'confirm','token':str(uuid.uuid4()),f'qty_{l.pk}':100,'date':'2026-10-23','address':'Form Address','carrier_type':'own','carrier':'Form Carrier','freight':'0','freight_currency':'CNY','notes':'Test'})
        self.assertEqual(response.status_code,302)
        self.assertTrue(summary(o)['complete'])
        response=self.client.get(f'/orders/{o.pk}/export/')
        data=b''.join(response.streaming_content)
        with ZipFile(BytesIO(data)) as z: self.assertIn('发货表.xlsx',z.namelist())

    def test_forged_owner_on_new_order_is_ignored(self):
        self.client.force_login(self.alice)
        self.client.post('/orders/new/',{'number':'FORGED','customer':'Test','owner':self.bob.pk,'date':'2026-10-23'})
        self.assertEqual(Order.objects.get(number='FORGED').owner_id,self.alice.pk)

    def test_duplicate_transaction_form_does_not_create_second(self):
        self.client.force_login(self.admin)
        data={'date':'2026-10-23','amount':'100','currency':'USD','party':'Company','platform':'Bank','fx':'1','fx_source':'Manual','purpose':'Test','reference':'DUPLICATE'}
        url=f'/orders/{self.order.pk}/transactions/receipt/new/'
        self.assertEqual(self.client.post(url,data).status_code,302)
        self.assertEqual(self.client.post(url,data).status_code,200)
        self.assertEqual(self.order.transactions.filter(reference='DUPLICATE').count(),1)
