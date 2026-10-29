
from pathlib import Path
from openpyxl import load_workbook, Workbook
root=Path('test-artifacts')
book=load_workbook(root/'product-template.xlsx')
assert book['产品录入'].max_column==9
assert book['产品录入']['A1'].value=='产品名称'
safe=load_workbook(root/'text-safety.xlsx')
assert safe.active['A2'].data_type=='s'
assert safe.active['A2'].value.startswith('=HYPERLINK')
assert safe.active['B2'].value==3
book['产品录入'].append(['虚构导入产品','IMPORT',5,8,'USD','自供货','',2,0])
book.save(root/'import-products.xlsx')
print('XLSX files verified using openpyxl; formula text is not executable.')
